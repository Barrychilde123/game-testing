# -*- coding: utf-8 -*-
"""生成 2048 游戏测试用例 Excel。"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "02-test-cases.xlsx")
HEADERS = ["用例编号", "模块", "用例标题", "前置条件", "操作步骤", "预期结果", "优先级", "测试类型", "执行结果", "备注"]

CASES = [
    ("GAME-001", "初始化", "新开游戏生成 2 个方块", "打开游戏页面", "1. 打开页面\n2. 读取棋盘状态", "棋盘 16 格，恰好 2 个非零方块，分数为 0，状态 playing", "P0", "功能", "通过", "自动化"),
    ("GAME-002", "初始化", "新游戏重置棋盘和分数", "已有游戏进度", "1. 点击“新游戏”\n2. 读取状态", "棋盘重置为 2 个方块，分数 0，状态 playing", "P0", "功能", "通过", "自动化"),
    ("GAME-003", "最高分", "新游戏保留最高分", "最高分为 104", "1. 点击“新游戏”\n2. 查看最高分", "最高分仍为 104，分数归 0", "P1", "功能", "通过", "自动化"),
    ("GAME-004", "合并", "左移合并 2+2", "棋盘第一行为 2,2,0,0", "1. 按左方向键\n2. 读取状态", "出现 4，分数 +4，生成 1 个新方块", "P0", "功能", "通过", "自动化"),
    ("GAME-005", "合并", "右移合并 2+2", "棋盘第一行为 0,0,2,2", "1. 按右方向键\n2. 读取状态", "出现 4，分数 +4", "P1", "功能", "通过", "自动化"),
    ("GAME-006", "合并", "上移合并 2+2", "第一列前两格为 2", "1. 按上方向键\n2. 读取状态", "出现 4，分数 +4", "P1", "功能", "通过", "自动化"),
    ("GAME-007", "合并", "下移合并 2+2", "第一列中间两格为 2", "1. 按下方向键\n2. 读取状态", "出现 4，分数 +4", "P1", "功能", "通过", "自动化"),
    ("GAME-008", "合并", "四个 2 合并为两个 4", "棋盘第一行为 2,2,2,2", "1. 按左方向键\n2. 读取状态", "出现两个 4，分数 +8", "P0", "功能", "通过", "自动化"),
    ("GAME-009", "边界", "无法移动时不生成新方块", "棋盘第一行为 2,4,0,0", "1. 按左方向键\n2. 读取状态", "move 返回 false，棋盘不变，分数不变", "P1", "边界", "通过", "自动化"),
    ("GAME-010", "胜负", "棋盘填满且无合并时游戏结束", "棋盘填满且相邻无相同", "1. 读取状态\n2. 查看弹窗", "状态 over，弹窗标题“游戏结束”", "P0", "功能", "通过", "自动化"),
    ("GAME-011", "胜负", "合成 2048 弹出胜利提示", "棋盘有 1024,1024", "1. 左移合并\n2. 查看状态", "出现 2048，状态 won，弹窗“你赢了！”", "P0", "功能", "通过", "自动化"),
    ("GAME-012", "胜负", "胜利后可以继续游戏", "已弹出胜利弹窗", "1. 点击“继续游戏”\n2. 查看状态", "弹窗关闭，状态回到 playing，可继续移动", "P1", "功能", "通过", "自动化"),
    ("GAME-013", "最高分", "最高分持久化", "得分超过历史最高分", "1. 刷新页面\n2. 查看最高分", "最高分从 localStorage 恢复，刷新后仍保留", "P1", "功能", "通过", "自动化"),
    ("GAME-014", "键盘", "方向键控制", "棋盘第一行为 2,2", "1. 发送 ArrowLeft\n2. 查看状态", "方块合并，分数 +4", "P1", "功能", "通过", "自动化"),
    ("GAME-015", "键盘", "WASD 控制", "棋盘第一行为 2,2", "1. 发送 a 键\n2. 查看状态", "方块合并，分数 +4", "P2", "功能", "通过", "自动化"),
    ("GAME-016", "兼容性", "375px 移动端不横向溢出", "移动视口 375x812", "1. 打开页面\n2. 读取 scrollWidth 与 board 宽度", "scrollWidth ≤ innerWidth，棋盘不溢出", "P1", "兼容性", "通过", "自动化"),
    ("GAME-017", "兼容性", "320px 小屏不横向溢出", "移动视口 320x568", "1. 打开页面\n2. 读取 scrollWidth 与 board 宽度", "scrollWidth ≤ innerWidth，棋盘不溢出", "P2", "兼容性", "通过", "自动化"),
    ("GAME-018", "性能", "弱网 3G 下可加载", "CDP 模拟 300ms 延迟、500KB/s", "1. 开启弱网\n2. 刷新页面\n3. 等待棋盘", "页面可加载，出现 2 个方块，加载 < 10s", "P1", "性能", "通过", "实测约 0.7s"),
    ("GAME-019", "稳定性", "连续 200 次操作不报错", "游戏运行中", "1. 循环执行 200 次方向操作\n2. 读取状态", "无 JS 异常，棋盘长度 16，状态合法", "P1", "稳定性", "通过", "实测约 2.3ms"),
    ("GAME-020", "性能", "正常网络加载 < 2s", "本地服务", "1. 打开页面\n2. 读取 navigation.loadEventEnd", "loadEventEnd < 2000ms", "P2", "性能", "通过", "实测约 69ms"),
    ("GAME-021", "兼容性", "16384 高数值方块样式", "棋盘有 8192,8192", "1. 合并生成 16384\n2. 查看方块样式", "16384 方块有深色背景和白色文字，可读", "P2", "兼容性", "通过", "自动化"),
    ("GAME-022", "异常", "localStorage 不可用时仍可游戏", "浏览器禁用 localStorage", "1. 注入 getItem/setItem 抛异常\n2. 打开页面", "页面正常加载，状态 playing，最高分 0", "P2", "异常", "通过", "自动化"),
]

wb = Workbook()
ws = wb.active
ws.title = "测试用例"
header_fill = PatternFill("solid", fgColor="1F5FA8")
header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
body_font = Font(name="Microsoft YaHei", size=9)
thin = Side(style="thin", color="C9D3E0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col, h in enumerate(HEADERS, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

for r, case in enumerate(CASES, 2):
    for c, value in enumerate(case, 1):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = body_font
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if c in (1, 7, 8, 9) else "left")

for col, w in {"A": 11, "B": 9, "C": 30, "D": 22, "E": 30, "F": 38, "G": 7, "H": 9, "I": 10, "J": 16}.items():
    ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 26
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:J{len(CASES)+1}"

dv = DataValidation(type="list", formula1='"未执行,通过,失败,阻塞,跳过"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"I2:I{len(CASES)+1}")
ws.conditional_formatting.add(f"I2:I{len(CASES)+1}", CellIsRule(operator="equal", formula=['"通过"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add(f"I2:I{len(CASES)+1}", CellIsRule(operator="equal", formula=['"失败"'], fill=PatternFill("solid", fgColor="FFC7CE")))

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved:", OUT, "cases:", len(CASES))
